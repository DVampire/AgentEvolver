"""Draw the original logical architecture layout with current component contracts.

Native PowerPoint text, shapes and connectors; no raster illustration. English and
Chinese share geometry. Export with the office/PDF pipeline in README.md.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parents[1] / "assets"
C = dict(bg="07100E", card="0D1B18", blue_bg="0C1D26", green_bg="0B241D",
         purple_bg="1C192C", amber_bg="252117", text="ECF7F2", muted="AEC4BA",
         blue="82B8D9", green="63E6B5", purple="B2A4DD", amber="F3BD71", border="38544A")


def build(lang):
    zh = lang == "zh"
    t = lambda en, cn: cn if zh else en
    p = Presentation()
    p.slide_width, p.slide_height = Inches(24), Inches(13.5)
    p.core_properties.title = t("Self-Evolving Agent System Architecture", "自进化 Agent System 架构图")
    p.core_properties.author = "AgentEvolver contributors"
    p.core_properties.subject = "Task orchestration, runtime communication, capability ecosystem and self-evolution"
    s = p.slides.add_slide(p.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(C["bg"])
    font = "Noto Sans CJK SC" if zh else "DejaVu Sans"

    def text(x, y, w, h, value, size=15, color="text", bold=False, center=False):
        shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = value.replace("\n", " / ")[:100]
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(value.split("\n")):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.name = font
            para.font.size = Pt(size)
            para.font.bold = bold
            para.font.color.rgb = RGBColor.from_string(C[color])
            para.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
            para.space_after = Pt(2)
            para.line_spacing = 1.06
            ea = OxmlElement("a:ea")
            ea.set("typeface", font)
            para._p.get_or_add_pPr().get_or_add_defRPr().append(ea)
        return shape

    def box(x, y, w, h, fill="card", border="border", radius=.10, kind=MSO_SHAPE.ROUNDED_RECTANGLE, width=1):
        sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = radius
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(C[fill])
        else:
            sh.fill.background()
        sh.line.color.rgb = RGBColor.from_string(C[border])
        sh.line.width = Pt(width)
        return sh

    def line(points, color="blue", arrow=False, both=False, dashed=False, width=1.6):
        for i, (a, b) in enumerate(zip(points, points[1:])):
            sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(a[0]), Inches(a[1]), Inches(b[0]), Inches(b[1]))
            sh.line.color.rgb = RGBColor.from_string(C[color])
            sh.line.width = Pt(width)
            ln = sh._element.spPr.get_or_add_ln()
            if dashed:
                dash = OxmlElement("a:prstDash"); dash.set("val", "dash"); ln.append(dash)
            if arrow and i == len(points)-2:
                end = OxmlElement("a:tailEnd"); end.set("type", "triangle"); ln.append(end)
            if both and i == 0:
                end = OxmlElement("a:headEnd"); end.set("type", "triangle"); ln.append(end)

    def icon(kind, x, y, size=.35, color="green"):
        # Icons are PowerPoint primitives, so they remain editable in the deck.
        def path(points):
            line([(x+u*size, y+v*size) for u,v in points], color, width=1.3)
        def oval(u,v,w,h):
            box(x+u*size,y+v*size,w*size,h*size,None,color,kind=MSO_SHAPE.OVAL,width=1.2)
        def frame(u=.08,v=.12,w=.84,h=.74):
            box(x+u*size,y+v*size,w*size,h*size,None,color,radius=.03,width=1.2)
        if kind == "user":
            oval(.34,.02,.32,.32); path([(.13,.96),(.13,.67),(.28,.49),(.72,.49),(.87,.67),(.87,.96)])
        elif kind == "code":
            path([(.31,.19),(.07,.5),(.31,.81)]);path([(.69,.19),(.93,.5),(.69,.81)]);path([(.58,.1),(.42,.9)])
        elif kind == "globe":
            oval(.05,.05,.9,.9);oval(.3,.05,.4,.9);path([(.06,.5),(.94,.5)]);path([(.16,.27),(.84,.27)]);path([(.16,.73),(.84,.73)])
        elif kind == "db":
            box(x+.1*size,y+.05*size,.8*size,.9*size,None,color,kind=MSO_SHAPE.CAN,width=1.2)
        elif kind == "shield":
            path([(.5,.03),(.9,.2),(.84,.65),(.5,.97),(.16,.65),(.1,.2),(.5,.03)])
            path([(.3,.46),(.46,.62),(.72,.34)])
        elif kind == "chart":
            path([(.06,.04),(.06,.94),(.96,.94)]);path([(.25,.77),(.25,.52)]);path([(.5,.77),(.5,.33)]);path([(.75,.77),(.75,.11)])
        elif kind == "monitor":
            frame(.06,.07,.88,.62);path([(.5,.69),(.5,.95)]);path([(.28,.95),(.72,.95)])
            path([(.15,.43),(.32,.43),(.4,.23),(.56,.55),(.65,.37),(.85,.37)])
        elif kind == "nodes":
            for u,v in [(.13,.16),(.72,.16),(.43,.73)]:oval(u,v,.2,.2)
            path([(.33,.26),(.72,.26)]);path([(.25,.36),(.5,.73)]);path([(.79,.36),(.58,.73)])
        elif kind == "terminal":
            frame();path([(.23,.31),(.4,.48),(.23,.65)]);path([(.53,.65),(.78,.65)])
        elif kind == "chat":
            path([(.07,.08),(.93,.08),(.93,.7),(.46,.7),(.22,.95),(.22,.7),(.07,.7),(.07,.08)])
            for u in [.29,.49,.69]:oval(u,.34,.04,.04)
        elif kind == "cube":
            path([(.5,.03),(.94,.26),(.94,.73),(.5,.97),(.06,.73),(.06,.26),(.5,.03),(.5,.5),(.94,.26)])
            path([(.06,.26),(.5,.5),(.5,.97)])
        elif kind == "loop":
            path([(.2,.78),(.07,.51),(.18,.21),(.48,.08),(.78,.2),(.88,.41)])
            path([(.64,.38),(.88,.41),(.91,.14)])
            path([(.8,.64),(.61,.88),(.3,.9)]);path([(.31,.69),(.3,.9),(.53,.98)])
        else:
            frame(.17,.03,.68,.93);path([(.3,.28),(.7,.28)]);path([(.3,.48),(.7,.48)]);path([(.3,.68),(.61,.68)])

    def panel(x,y,w,h,title,color,fill,size=20):
        box(x,y,w,h,fill,color,radius=.05,width=1.5)
        text(x+.18,y+.16,w-.36,.38,title,size,"muted" if color == "border" else color,True,True)

    def chip(x,y,w,h,label,color="blue",symbol=None,size=14):
        box(x,y,w,h,"card",color,radius=.13,width=.8)
        if symbol:
            icon(symbol,x+.12,y+(h-.32)/2,.32,color)
            text(x+.55,y+.12,w-.65,h-.15,label,size)
        else:
            text(x+.06,y+.12,w-.12,h-.15,label,size,center=True)

    # Same spatial anchors as the original diagram: task / actors / runtime /
    # ecosystem / evolution / versions / external world / infrastructure.
    text(.5,.22,23,.60,t("Self-Evolving Agent System Architecture", "自进化 Agent System 架构图"),32,bold=True,center=True)
    text(.5,.94,23,.30,t("Protocol communication  ·  Capability ecosystem  ·  Self-evolution loop", "协议通信  ·  能力生态  ·  自进化闭环"),16,"muted",center=True)

    # Top task path.
    box(.4,1.48,1.8,1.02,"blue_bg","blue")
    icon("user",.62,1.72,.48,"blue")
    text(1.2,1.83,.9,.35,t("User","用户"),17,bold=True)
    box(2.85,1.48,2.67,1.02,"blue_bg","blue")
    icon("doc",3.03,1.70,.48,"blue")
    text(3.64,1.68,1.75,.35,t("Task request","任务输入"),17,bold=True)
    text(3.64,2.13,1.72,.22,t("text · files · manifest","正文 · 附件 · manifest"),9.5,"muted")
    line([(2.2,1.99),(2.85,1.99)],arrow=True)
    line([(5.52,1.99),(6.12,1.99)],arrow=True)
    panel(6.12,1.42,7.45,1.66,t("MetaAgent / Builder Orchestrator","MetaAgent / Builder 总控编排器"),"blue","blue_bg",21)
    steps=[t("Plan","规划"),t("Decompose","分解"),t("Delegate*","委派*"),t("Monitor","监控"),t("Conclude","汇总")]
    for i,label in enumerate(steps):chip(6.35+i*1.40,2.05,1.25,.5,label,size=11.5)
    text(6.35,2.70,6.96,.24,t("Shared Agent loop · Website / Game Builder can work solo", "共用 Agent 循环 · Website / Game Builder 可独立工作"),10.5,"muted",center=True)

    # Left-side actor family: same job and position as the original figure.
    panel(.45,3.30,3.46,4.02,t("Actor Agents","执行型 Agent"),"green","green_bg",21)
    for i,(name,cn,kind) in enumerate([
        ("Code Agent","代码 Agent","code"),("General Agent","通用 Agent","chat"),
        ("Browser Agent","浏览器 Agent","globe"),("Computer Agent","桌面 Agent","monitor"),
        ("Monitor Agent","监控 Agent","chart"),("Reviewer Agent","审查 Agent","shield")]):
        chip(.68,3.98+i*.45,3.0,.37,t(name,cn),"green",kind,size=12.4)
    text(.63,6.89,3.1,.27,t("* Delegation when enabled", "* 按配置和权限启用委派"),10.3,"muted",center=True)

    # Central runtime and communication block, retaining the original two halves.
    panel(4.80,3.92,9.80,2.95,t("Agent Runtime & Protocol","Agent Runtime 与通信协议"),"blue","blue_bg",22)
    box(5.02,4.55,4.49,1.55,"card","blue",radius=.04)
    box(9.75,4.55,4.61,1.55,"card","purple",radius=.04)
    text(5.17,4.67,4.18,.33,t("Runtime kernel","Runtime 进程内核"),16,"blue",True,True)
    text(9.92,4.67,4.26,.33,t("Messages & interactions","消息与交互语义"),16,"purple",True,True)
    for x,y,label,kind in [
        (5.22,5.16,t("Mailbox / signals","邮箱 / 信号"),"chat"),
        (7.46,5.16,t("Spawn / lifecycle","创建 / 生命周期"),"nodes"),
        (5.22,5.66,t("Suspend / resume","暂停 / 恢复"),"loop"),
        (7.46,5.66,t("Shared budget","共享运行预算"),"shield")]:
        icon(kind,x,y,.24,"blue");text(x+.36,y,1.86,.32,label,10.8)
    for x,y,label,kind in [
        (9.95,5.16,t("Dispatch / reply","委派 / 回复"),"nodes"),
        (12.19,5.16,t("Progress / control","进度 / 控制"),"monitor"),
        (9.95,5.66,t("Publish / subscribe","发布 / 订阅"),"chat"),
        (12.19,5.66,t("Escalation / events","上报 / 事件"),"doc")]:
        icon(kind,x,y,.24,"purple");text(x+.36,y,1.82,.32,label,10.5)
    box(5.02,6.28,9.34,.37,"card","border",radius=.1)
    text(5.18,6.36,9.02,.22,t("Agent loop: ContextAssembler → ModelManager → ToolRouter → observations", "Agent 循环：ContextAssembler → ModelManager → ToolRouter → 环境观察"),10.6,"muted",center=True)
    text(6.00,7.03,8.4,.27,t("Agent communication and lifecycle pass through the runtime kernel", "Agent 间通信及生命周期由 Runtime 内核统一管理"),11,"muted",center=True)

    # Orthogonal channels between controller, actors and runtime.
    line([(7.15,3.08),(7.15,3.49),(6.10,3.49),(6.10,3.92)],arrow=True)
    line([(9.88,3.92),(9.88,3.08)],arrow=True,both=True)
    line([(12.68,3.08),(12.68,3.49),(13.60,3.49),(13.60,3.92)],arrow=True)
    line([(3.91,4.16),(4.32,4.16),(4.32,6.60),(3.91,6.60)],arrow=True,both=True,dashed=True)
    line([(4.32,5.52),(4.80,5.52)],arrow=True,both=True)
    text(6.30,3.19,2.24,.24,t("tasks / results","任务 / 结果"),10,"blue")

    # Self-evolution retains the original ring and the three supporting boxes.
    panel(15.35,1.42,8.20,5.22,t("Self-Evolution Engine","自进化引擎 / Self-Evolution Engine"),"purple","purple_bg",22)
    nodes=[
        (16.87,2.10,2.23,.55,"1",t("Observe / probe","发现 / 探测")),
        (18.66,3.00,1.98,.63,"2",t("Build / register","生成 / 注册")),
        (18.32,4.39,2.32,.63,"3",t("Evaluate","评估")),
        (15.98,4.39,2.06,.63,"4",t("Keep / roll back","采纳 / 回滚")),
        (15.65,3.00,2.24,.63,"5",t("Use / revisit","使用 / 再审视")),
    ]
    for x,y,w,h,n,label in nodes:
        box(x,y,w,h,"card","purple",radius=.3)
        box(x+.12,y+.15,.3,.3,"purple_bg","purple",kind=MSO_SHAPE.OVAL)
        text(x+.12,y+.195,.3,.2,n,10.5,"purple",True,True)
        text(x+.51,y+.20,w-.60,.29,label,11.4,bold=True,center=True)
    line([(19.10,2.37),(19.65,2.37),(19.65,3.0)],"purple",True)
    line([(19.65,3.63),(19.65,4.39)],"purple",True)
    line([(18.32,4.70),(18.04,4.70)],"purple",True)
    line([(16.48,4.39),(16.48,3.63)],"purple",True)
    line([(16.75,3.0),(16.75,2.37),(16.87,2.37)],"purple",True)
    icon("loop",17.97,3.33,.55,"purple")
    text(17.69,4.03,1.10,.24,t("iterate","持续迭代"),10.5,"purple",center=True)
    for x,y,w,h,title,detail,kind in [
        (21.00,2.15,2.27,.79,"self_evolving_skill",t("Author / refine methods","编写 / 改进方法"),"doc"),
        (21.00,3.31,2.27,.79,t("Agent self-review","Agent 自行评估"),t("Baseline + reuse + cost","基线 + 复用 + 成本"),"shield"),
        (21.00,4.47,2.27,.79,"adoption_tool",t("Version / decision / use","版本 / 决策 / 使用"),"nodes")]:
        box(x,y,w,h,"card","purple",radius=.07)
        icon(kind,x+.12,y+.18,.3,"purple")
        text(x+.50,y+.15,w-.58,.24,title,10.0,bold=True)
        text(x+.12,y+.51,w-.24,.23,detail,9.4,"muted",center=True)
    line([(20.64,3.15),(20.82,3.15),(20.82,2.55),(21.00,2.55)],"purple",True,dashed=True,width=1.0)
    line([(20.64,4.61),(20.82,4.61),(20.82,3.71),(21.00,3.71)],"purple",True,dashed=True,width=1.0)
    line([(17.00,5.02),(17.00,5.39),(20.82,5.39),(20.82,4.87),(21.00,4.87)],"purple",True,dashed=True,width=1.0)
    box(15.70,5.65,3.30,.48,"card","purple",radius=.1)
    icon("shield",15.84,5.74,.28,"purple")
    text(16.22,5.79,2.60,.22,"scope + enable_evolving",10.5,bold=True)
    text(19.27,5.66,3.99,.57,t("Version-bound call evidence\nSelf-review by the acting Agent", "调用证据绑定确切版本\n由执行任务的 Agent 自评"),10.5,"muted",center=True)
    line([(13.57,2.22),(14.82,2.22),(14.82,2.66),(15.35,2.66)],"purple",True)
    text(13.78,1.57,1.44,.59,t("Observed need /\ncapability gap", "真实需求 /\n能力改进机会"),10.2,"purple",center=True)
    line([(14.60,5.19),(15.35,5.19)],"purple",True,both=True)

    # Capability ecosystem: same broad horizontal layer, with the two missing
    # extension families made explicit rather than hidden in another module.
    panel(1.15,7.85,13.45,2.15,t("Evolvable Capability Ecosystem","可进化能力生态 / Capability Ecosystem"),"green","green_bg",21)
    caps=[
        ("Agent + Prompt","Agent + Prompt","nodes",t("Behavior / roles","行为 / 角色")),
        ("Skills","技能 Skill","doc",t("Methods / scripts","方法 / 脚本")),
        ("Tools","工具 Tool","terminal",t("Bash / inspect","Bash / 检查")),
        ("Workflows","工作流","nodes",t("Compose / run","组合 / 执行")),
        ("Environments","环境","cube",t("Browser / Godot","浏览器 / Godot")),
        ("Memory","记忆","db",t("Notes / retrieval","笔记 / 检索")),
        ("Connectors","连接器","globe",t("MCP / APIs","MCP / API")),
        ("Plugins","插件","cube",t("Integrations","能力集成")),
    ]
    for i,(en,cn,kind,detail) in enumerate(caps):
        x=1.34+i*1.64
        box(x,8.50,1.55,1.25,"card","green",radius=.06,width=.8)
        icon(kind,x+.57,8.64,.38,"green")
        text(x+.04,9.14,1.47,.28,t(en,cn),10.8,bold=True,center=True)
        text(x+.04,9.51,1.47,.22,detail,9.1,"muted",center=True)
    line([(2.10,7.32),(2.10,7.85)],"green",True,both=True)
    line([(5.40,6.87),(5.40,7.55),(5.40,7.85)],"green",True,both=True)
    line([(2.10,7.55),(14.93,7.55)],"green",dashed=True,width=1.1)

    # Extension manager stays immediately below evolution, with the original
    # archive/manifest/load/rollback structure and a path back to capabilities.
    panel(15.25,7.19,4.85,2.81,t("Extension Manager & Versioning","扩展管理与版本控制"),"amber","amber_bg",16.5)
    icon("db",15.47,8.10,.58,"amber")
    for x,y,label,kind in [
        (16.22,7.92,t("Active versions","活动版本"),"nodes"),
        (16.22,8.50,t("Version archive","版本归档"),"db"),
        (16.22,9.08,"manifest.json","doc")]:
        chip(x,y,2.07,.43,label,"amber",kind,10.1)
    chip(18.45,7.92,1.43,.62,t("Admission /\nload","准入 / 加载"),"amber",size=10.1)
    chip(18.45,8.76,1.43,.62,t("Rollback /\nunload","回滚 / 卸载"),"amber",size=10.1)
    text(15.52,9.70,4.30,.20,t("extension/ · outside the framework core","extension/ · 位于框架核心之外"),9.0,"muted",center=True)
    line([(17.73,6.64),(17.73,7.19)],"amber",True)
    line([(23.27,4.87),(23.40,4.87),(23.40,6.86),(19.23,6.86),(19.23,7.19)],"amber",True)
    line([(15.25,8.58),(14.93,8.58),(14.93,7.55)],"green",dashed=True,width=1.3)
    line([(14.93,8.58),(14.60,8.58)],"green",True,dashed=True)

    # Far-right external systems, with published products and native Godot added.
    panel(20.75,7.19,2.80,3.88,t("External world","外部世界"),"blue","blue_bg",17)
    external=[
        (t("Web / websites","Web / 网站"),"globe"),
        (t("APIs / services","API / 服务"),"nodes"),
        (t("Data / databases","数据 / 数据库"),"db"),
        (t("MCP servers","MCP 服务"),"cube"),
        (t("Browser / Godot","浏览器 / Godot"),"monitor"),
        (t("Published products","已发布产品"),"globe")]
    for i,(label,kind) in enumerate(external):chip(20.94,7.85+i*.48,2.41,.38,label,"blue",kind,10.4)
    # Capability traffic takes its own rail; it does not pass through the version manager.
    line([(14.60,9.90),(14.93,9.90),(14.93,10.20),(20.42,10.20),(20.42,8.03),(20.75,8.03)],"blue",True)
    line([(20.42,8.51),(20.75,8.51)],"blue",True)
    line([(20.42,9.00),(20.75,9.00)],"blue",True)
    line([(20.42,9.48),(20.75,9.48)],"blue",True)
    line([(20.42,9.97),(20.75,9.97)],"blue",True)
    line([(20.42,10.20),(20.42,10.46),(20.75,10.46)],"blue",True)

    # Supporting infrastructure retains the bottom band. Plan/context and
    # deployment are additions within it, not replacements for the main graph.
    panel(.45,10.55,19.55,2.08,t("Supporting Infrastructure","支撑基础设施 / Supporting Infrastructure"),"blue","blue_bg",21)
    infra=[
        (t("Trace / views","Trace / 观测"),"monitor",t("Events / monitors","事件 / 监控视图")),
        ("Trajectory","doc",t("SFT / RL export","SFT / RL 导出")),
        ("Registry","nodes",t("Discover / register","发现 / 注册")),
        (t("Task / Session","任务 / 会话"),"doc",t("Identity / workspace","身份 / 工作空间")),
        ("Plan / Context","doc",t("index.md / plan.md\nLayers / compaction","index.md / plan.md\n分层 / 压缩")),
        (t("Permissions","权限 / 约束"),"shield",t("Policy / budgets","策略 / 预算")),
        ("Config / Version","nodes",t("Assembly / lineage","配置 / 版本关系")),
        (t("Model Manager","模型管理"),"nodes",t("Routing / usage","路由 / 用量")),
        ("Deploy / Gateway","globe",t("URLs / product access","URL / 产品访问")),
        ("Benchmark / Data","chart",t("Tasks / evaluation","任务 / 评估")),
    ]
    for i,(label,kind,detail) in enumerate(infra):
        x=.67+i*1.93
        box(x,11.20,1.82,1.21,"card","blue",radius=.05,width=.8)
        icon(kind,x+.73,11.32,.34,"blue")
        text(x+.03,11.79,1.76,.25,label,10.1,bold=True,center=True)
        text(x+.03,12.05,1.76,.35,detail,8.6,"muted",center=True)
    for x in [2.10,5.39,8.68,11.97,14.01]:line([(x,10.55),(x,10.00)],"green",True,both=True,width=1.0)
    # Trace/history feedback and benchmark evidence follow the perimeter like the original.
    line([(1.10,12.63),(1.10,12.99),(8.16,12.99),(8.16,12.63)],"blue",True,dashed=True,width=1.0)
    text(2.10,13.08,10.7,.22,t("Trace and retained context preserve evidence for later decisions", "Trace 与持久上下文为后续决策保留可追溯证据"),10.0,"muted")
    line([(19.00,12.63),(19.00,13.31),(23.81,13.31),(23.81,3.70),(23.55,3.70)],"blue",True,dashed=True,width=1.0)
    text(13.02,12.90,5.62,.22,t("Benchmark / task checks → evaluation evidence", "Benchmark / 任务检查 → 评估证据"),9.8,"muted")

    # The original bottom-right legend makes the different relationships explicit.
    panel(20.75,11.43,2.80,1.62,t("Legend","图例"),"border","card",13.5)
    for i,(label,color,dash) in enumerate([
        (t("Task / communication","任务 / 通信"),"blue",False),
        (t("Capability use","能力调用"),"green",False),
        (t("Evolution / adoption","进化 / 采纳"),"purple",False),
        (t("Extension lifecycle","扩展生命周期"),"amber",False),
        (t("Evidence / feedback","证据 / 反馈"),"blue",True)]):
        y=11.97+i*.19
        line([(20.93,y+.06),(21.31,y+.06)],color,True,dashed=dash,width=1.1)
        text(21.42,y,1.91,.2,label,8.7,"muted")

    s.notes_slide.notes_text_frame.text = (
        "Preserves the original architecture's spatial organization: task/orchestrator at top, "
        "actors at left, runtime and protocol in the center, capability ecosystem beneath, "
        "self-evolution cycle and version management at right, external world on the far right, "
        "supporting infrastructure at bottom. Source map: docs/diagrams/README.md. "
        "Protocol labels describe runtime message contracts, not a separate protocol package. "
        "Builder is an alternative root role; actor delegation is optional. Evolution is performed "
        "by the acting Agent with self_evolving_skill and adoption_tool, not mandatory generator/evaluator agents. "
        "Use after keep verifies the consumer; rollback returns to investigation instead. "
        "All shapes and icons are native editable PowerPoint objects."
    )
    # Disable theme shadow effects, which office export can rasterize.
    for shape in s.shapes:
        shape._element.spPr.append(OxmlElement("a:effectLst"))
        for effect in shape._element.xpath(".//a:effectRef"):
            effect.set("idx", "0")
    p.save(OUT / ("arch_zh.pptx" if zh else "arch.pptx"))


if __name__ == "__main__":
    for language in ("en", "zh"):
        build(language)
